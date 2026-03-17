# Standard library imports
import logging
import os
import time
from pathlib import Path
from typing import Any, Dict

# Third-party imports
import pandas as pd  # type: ignore[reportMissingTypeStubs]
import pytesseract  # type: ignore[reportMissingTypeStubs]
from chromadb.api import Collection  # type: ignore[reportPrivateImportUsage]
from langdetect import detect  # type: ignore[reportMissingTypeStubs]
from openpyxl.worksheet.worksheet import Worksheet
from pdf2image import \
    convert_from_path  # type: ignore[reportUnknownVariableType]
from pdfminer.high_level import extract_text as extractPDF
from PIL import Image

from Algos.Masker import Masker
from Algos.UnicodeNormalizer import UnicodeNormalizer
from Commons.Exceptions import DataProcessingError, DocumentsDirError
from Commons.SingletonMixin import SingletonMixin
from Compliance.Exclusions import Exclusions
from Config.Config import Config
from Globals.CounterInstance import (ExclusionsCount, FailedCount,
                                     IgnoredCount, ProcessedCount)
from Globals.Globals import Globals
from Gui.Colors import BRIGHT_BLUE, CYAN, ORANGE, RESET
from Gui.PrettyWriter import PrettyWriter
from Helpers.ChromaDBHelper import ChromaDBHelper
from Helpers.CSVWriter import CSVWriter
from Helpers.FileUtils import FileUtils
from Helpers.Helpers import Helpers
from Helpers.OfficeDocConverter import OfficeDocConverter
from Helpers.ValidExtensions import ValidExtensions
from Strategies.ProcessingStrategy import ProcessingStrategy
from Strategies.StrategyType import StrategyType


class LoadAndClassifyProcessor(SingletonMixin):
    """
    Singleton that walks a work directory, extracts text from many file
    formats (PDF, Office, images, etc.), and delegates either to a
    home-brew chunker or to the classification implementation.
    """

    def process(self) -> None:
        self.process_files()

    def __init__(
        self,
        strategy: ProcessingStrategy,
        *,
        cfg: "Config | None" = None,
        pretty: "PrettyWriter | None" = None,
        helpers: "Helpers | None" = None,
    ) -> None:
        if self._initialized:
            return
        self._initialized = True
        self._strategy: ProcessingStrategy = strategy

        # Utilities & counters
        self._failed_countInstance: FailedCount = FailedCount()
        self._processed_countInstance: ProcessedCount = ProcessedCount()
        self._ignored_countInstance: IgnoredCount = IgnoredCount()
        self._exclusions_countInstance: ExclusionsCount = ExclusionsCount()
        self.exclusions: Exclusions = Exclusions()
        self.pretty: PrettyWriter = pretty or PrettyWriter()
        self.helpers: Helpers = helpers or Helpers()
        self.fileUtils: FileUtils = FileUtils()
        self.csvWriter: CSVWriter = CSVWriter()
        self.masker: Masker = Masker()
        self.unicode_normalizer: UnicodeNormalizer = UnicodeNormalizer()

        self.fileName: str | None = None
        self.fileHash: str = "N/A"

        # Validators, converters & delegates
        self.valid_extsInstance: ValidExtensions = ValidExtensions()
        self.office_convInstance: OfficeDocConverter = OfficeDocConverter()
        self.globalsInstance: Globals = Globals()

        # Configuration
        self.cfg = cfg or Config()
        self.doc_dir: str = self.cfg.get_str("DOC_DIR")
        p = Path(self.doc_dir)
        # readable
        if not os.access(p, os.R_OK):
            self.pretty.write(
                "E",
                "Documents directory",
                f"Path {self.doc_dir} is not readable. Check DOC_DIR in Configuration/Config_Global.py",
            )
            raise DocumentsDirError

        self.use_exclusions: bool = self.cfg.get_bool("USE_EXCLUSIONS")
        self.process_unchanged: bool = self.cfg.get_bool("_PROCESS_IF_UNCHANGED")
        self.friendly_name: str = self.cfg.get_str("_FRIENDLY_NAME")
        self.collection: Collection | None = None
        self.consider_as_text_file: list[str] = self.cfg.get_list(
            "_CONSIDER_AS_TEXT_FILE"
        )
        self.office_doc_extraction: Dict[str, bool] = self.cfg.get_dict(
            "_OFFICE_DOC_EXTRACTION"
        )

        self.chromaDBHelper: ChromaDBHelper = ChromaDBHelper()
        if self.friendly_name == "RAGLoad":
            self.collection_name: str
            self.collection_name, _ = self.chromaDBHelper.change_chroma_collection(
                self.cfg.get("COLLECTION"), True  # type: ignore[reportArgumentType]
            )

        self.helpers.configure_tesseract()

    def worksheet_to_dataframe(self, ws: Worksheet) -> pd.DataFrame:
        rows: list[Any] = list(ws.iter_rows(values_only=True))
        if not rows:
            return pd.DataFrame()
        return pd.DataFrame(rows[1:], columns=rows[0])

    def docChanged(self) -> bool:
        # ————————————————
        # 1) OPEN CHROMA & GET STORED HASH
        # ————————————————
        self.coll_name, self.persist_directory = (
            self.chromaDBHelper.chroma_coll_name_and_mkdir_or_del(
                "create", self.collection_name
            )
        )
        self.client, self.collection = (
            self.chromaDBHelper.get_chroma_client_and_collection(
                self.persist_directory, self.coll_name, stamp=True
            )
        )

        # 2) ask for up to one matching entry’s metadata
        resp: Any = self.collection.get(
            where={"FilePath": self.escapedFilePath},
            limit=1,
            include=["metadatas", "documents"],
        )

        self.fileHash = self.fileUtils.hash_file(self.escapedFilePath)
        # 3) pull out FileHash if it exists
        metas: list[Any] = resp.get("metadatas", [])
        if metas:
            prev_hash: str | None = metas[0].get("FileHash")
        else:
            prev_hash = None

        # ————————————————
        # 2) COMPUTE CURRENT HASH & COMPARE
        # ————————————————
        if prev_hash == self.fileHash:
            if self.process_unchanged is True:
                self.pretty.write(
                    "I", "Vector store", f"No change in {self.escapedFilePath}"
                )
                self.pretty.write(
                    "I",
                    "Vector store",
                    f"but PROCESS_UNCHANGED {self.process_unchanged} override",
                )
                return True
            self.pretty.write(
                "A",
                "Vector Store",
                f"No change in {self.escapedFilePath}. Skip",
                color=CYAN,
            )
            self._processed_countInstance.increment()
            self.csvWriter.write_json2csv(
                {"FilePath": self.escapedFilePath, "Status": "NO CHANGE IN FILE"},
                "OK",
            )
            return False
        else:
            return True

    def _make_doc(self) -> dict[str, Any]:
        self.content, _ = self.helpers.safe_decode_to_unicode(
            self.content, True
        )  # Assuming not already UTF-8.
        detected: str = str(detect(self.content))  # type: ignore[reportUnknownArgumentType]
        language: str = detected
        self.doc = {
            "meta": {
                "FileName": self.fileName,
                "FilePath": self.escapedFilePath,
                "CreationDate": self.creation_date,
                "FileType": self.ftype,
                "Language": language,
                "WordCount": self.fileUtils.count_words(self.content),
                "FileHash": self.fileHash,
            },
            "content": self.content,
        }

        return self.doc

    def process_files(self) -> None:
        """
        Walk self.doc_dir, extract text from each file, then either
        chunk or classify.
        """
        for root, _, files in os.walk(self.doc_dir):
            for self.fileName in files:
                self.filePath: str = os.path.join(root, self.fileName)
                self.escapedFilePath: str = self.fileUtils.normalize_path(self.filePath)
                if self.use_exclusions and self.exclusions.contains(
                    self.escapedFilePath
                ):
                    self.pretty.write(
                        "W",
                        "EXCLUSIONS",
                        f"Excluding: {self.escapedFilePath}",
                        color=ORANGE,
                    )
                    self._exclusions_countInstance.increment()
                    continue

                self.pretty.write(
                    "I",
                    "START",
                    f"{BRIGHT_BLUE}Processing: {self.escapedFilePath}{RESET}",
                )

                # validate extension
                self.ftype: str = self.valid_extsInstance.getFileType(
                    self.escapedFilePath
                )
                if not self.valid_extsInstance.check(self.fileName, self.ftype):
                    self.pretty.write(
                        "I",
                        "Ignored extensions",
                        f"Ignored (invalid ext): {self.escapedFilePath} ({self.ftype})",
                    )
                    self._ignored_countInstance.increment()
                    continue

                if (
                    self._strategy.strategy_type == StrategyType.CHUNKS_TO_DB
                    and self.docChanged() == False
                ):
                    continue
                # fetch creation timestamp
                try:
                    c_ts: float = os.path.getctime(self.escapedFilePath)
                    self.creation_date: str = time.ctime(c_ts)
                except OSError:
                    self.creation_date = ""

                # extract text
                self.pretty.write("I", "Extract text", f"Extracting text from Document")
                self.content: str = ""
                disabled_office_component: str = ""
                try:
                    if self.valid_extsInstance.check(self.fileName, ["pdf"]):
                        logging.getLogger("pdfminer").setLevel(logging.ERROR)
                        self.content = extractPDF(self.escapedFilePath).strip()
                        if not self.content:
                            for page in convert_from_path(
                                self.escapedFilePath, dpi=300
                            ):
                                try:
                                    # Process the page
                                    self.content = (
                                        str(pytesseract.image_to_string(page)) + "\n"  # type: ignore[reportUnknownMemberType]
                                    )
                                finally:
                                    page.close()

                    # _CONSIDER_AS_TEXT_FILE: plain-text formats whose
                    # content is read as-is (txt, md, py, csv, log, …).
                    elif self.valid_extsInstance.check(
                        self.fileName, self.consider_as_text_file
                    ):
                        with open(self.escapedFilePath, "r", encoding="utf-8") as f:
                            self.content = f.read()

                    elif self.valid_extsInstance.check(self.fileName, ["doc", "docx"]):
                        if self.office_doc_extraction.get("Word"):
                            _, doc_obj = self.office_convInstance.convert_office_file(
                                self.escapedFilePath
                            )
                            self.content = "\n".join(p.text for p in doc_obj.paragraphs)
                            self.content += "\n".join(
                                cell.text
                                for tbl in doc_obj.tables
                                for row in tbl.rows
                                for cell in row.cells
                            )
                        else:
                            disabled_office_component = "MS Word"

                    elif self.valid_extsInstance.check(self.fileName, ["ppt", "pptx"]):
                        if self.office_doc_extraction.get("Power Point"):
                            _, pres = self.office_convInstance.convert_office_file(
                                self.escapedFilePath
                            )
                            for slide in pres.slides:
                                for shape in slide.shapes:
                                    if hasattr(shape, "text"):
                                        self.content += shape.text + "\n"
                        else:
                            disabled_office_component = "MS Power Point"

                    elif self.valid_extsInstance.check(self.fileName, ["xls", "xlsx"]):
                        if self.office_doc_extraction.get("Excel"):
                            _, wb = self.office_convInstance.convert_office_file(
                                self.escapedFilePath
                            )
                            df = self.worksheet_to_dataframe(wb.active)
                            self.content = str(df.to_string(index=False))  # type: ignore[reportUnknownMemberType]
                        else:
                            disabled_office_component = "MS Excel"

                    elif self.valid_extsInstance.check(
                        self.fileName,
                        ["png", "jpg", "jpeg", "gif", "bmp", "tiff", "webp"],
                    ):
                        img = Image.open(self.escapedFilePath)
                        if self.fileName.lower().endswith(".webp"):
                            self.pretty.write(
                                "I", "Conversion", "Converting WebP → RGB for OCR"
                            )
                            img = img.convert("RGB")
                        self.content = str(pytesseract.image_to_string(img))  # type: ignore[reportUnknownMemberType]

                except DataProcessingError:
                    raise
                except Exception as e:
                    self.pretty.write(
                        "W",
                        "Extraction fail",
                        f"Extraction failed: {self.escapedFilePath}: {e}",
                    )
                    # record a failed doc stub and continue
                    self._failed_countInstance.increment()
                    self.globalsInstance.add_failed_doc(
                        {"FilePath": self.escapedFilePath, "error": str(e)}
                    )
                    continue
                if self.content == "":
                    if disabled_office_component != "":
                        self.pretty.write(
                            "W",
                            "Office component",
                            f"File {self.escapedFilePath} was not processed because office component {disabled_office_component} is disabled in _OFFICE_DOC_EXTRACTION (Configuration/Config_Global.py)",
                        )
                    else:
                        self.pretty.write(
                            "I",
                            "Empty file",
                            f"File {self.escapedFilePath} has no content and is ignored",
                        )
                    continue
                self.content = self.unicode_normalizer.normalize(self.content)
                self.content = self.masker.mask(self.content)
                # build the document object
                self.doc = self._make_doc()

                # route to chunker or classifier
                self._strategy.process(self.doc)

                self._processed_countInstance.increment()
